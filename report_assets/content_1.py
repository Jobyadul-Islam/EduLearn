# -*- coding: utf-8 -*-
"""Front matter + Chapters 1-3."""
from docx_engine import *
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
import os

CERT_PATH = os.path.join(os.path.dirname(__file__), "..", "Practicum_Certificate.png")


def build(doc):
    # ---------------------------------------------------------------- Cover
    for _ in range(3):
        add_spacer(doc, 6)
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("EduLearn"); set_run_font(r, 24, bold=True)
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("A Web-Based E-Learning Platform for Course Delivery, Student\nProgress Tracking, and Instructor Management")
    set_run_font(r, 16, bold=True)
    add_spacer(doc, 30)
    for line in ["Md. Jobyadul Islam", "ID# 22303363"]:
        p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(line); set_run_font(r, 13)
    add_spacer(doc, 30)
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("A Practicum in the Partial Fulfillment of the Requirements\nfor the Award of Bachelor of Computer Science and Engineering (BCSE)")
    set_run_font(r, 12)
    add_spacer(doc, 40)
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("Department of Computer Science and Engineering\nCollege of Engineering and Technology\nIUBAT—International University of Business Agriculture and Technology")
    set_run_font(r, 12)
    add_spacer(doc, 20)
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("Summer 2026"); set_run_font(r, 12, bold=True)

    # -------------------------------------------------- Approval / Signature
    doc.add_page_break()
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("EduLearn — A Web-Based E-Learning Platform for Course Delivery,\nStudent Progress Tracking, and Instructor Management")
    set_run_font(r, 15, bold=True)
    add_spacer(doc, 10)
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("Md. Jobyadul Islam"); set_run_font(r, 12)
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("A Practicum in the Partial Fulfillment of the Requirements for the Award of Bachelor of\nComputer Science and Engineering (BCSE)")
    set_run_font(r, 11)
    add_spacer(doc, 10)
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("The practicum has been examined and approved,"); set_run_font(r, 12)
    add_spacer(doc, 26)
    for name, role in [("Prof. Dr. Utpal Kanti Das", "Chairman\nDept of CSE"),
                        ("Shahinur Alam", "Co-supervisor, Coordinator and Assistant Professor"),
                        ("Dr. Md. Abdul Awal", "Associate Professor\nSupervisor")]:
        p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run("_____________________________"); set_run_font(r, 12)
        p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(name); set_run_font(r, 12, bold=True)
        p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(role); set_run_font(r, 11)
        add_spacer(doc, 14)
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("Department of Computer Science and Engineering\nIUBAT—International University of Business Agriculture and Technology\nSummer 2026")
    set_run_font(r, 11)

    # -------------------------------------------------- Letter of Transmittal
    add_heading1(doc, "Letter of Transmittal")
    add_para(doc, "17 September 2026", align="left", indent_first=False)
    add_para(doc, "The Chair\nPracticum Defense Committee\nDepartment of Computer Science and Engineering\nIUBAT—International University of Business Agriculture and Technology\n4 Embankment Drive Road, Sector 10, Uttara Model Town\nDhaka 1230, Bangladesh.", align="left", indent_first=False)
    add_spacer(doc)
    p = doc.add_paragraph(); r = p.add_run("Subject: "); set_run_font(r, 12, bold=True)
    r2 = p.add_run("Letter of Transmittal."); set_run_font(r2, 12)
    r2.underline = True
    add_spacer(doc)
    add_para(doc, "Dear Sir,", align="left", indent_first=False)
    add_para(doc, "It is my pleasure to submit my practicum report entitled “EduLearn — A Web-Based E-Learning Platform for Course Delivery, Student Progress Tracking, and Instructor Management.”")
    add_para(doc, "This project has provided me with valuable practical experience in full-stack web application development, database design, role-based authentication and authorization, payment integration, automated testing, and building a real, working product from requirements through to a defensible, deployed system.")
    add_para(doc, "I sincerely thank the department and my respected faculty members, as well as Touch & Solve Technologies Ltd., for giving me the opportunity to complete this practicum project. I will be pleased to provide any further clarification if required.")
    add_para(doc, "Yours sincerely,", align="left", indent_first=False)
    add_spacer(doc, 20)
    add_para(doc, "_____________\n\nMd. Jobyadul Islam\n22303363", align="left", indent_first=False)

    # -------------------------------------------------- Organization's Certificate
    add_heading1(doc, "Organization's Certificate")
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    if os.path.exists(CERT_PATH):
        run = p.add_run()
        run.add_picture(CERT_PATH, width=Inches(5.8))
    add_spacer(doc, 8)
    add_para(doc, "The certificate above was issued by Touch & Solve Technologies Ltd. to Md. Jobyadul Islam on the successful completion of an internship in Web Development, undertaken as part of the practicum requirement for the Bachelor of Computer Science and Engineering (BCSE) program at IUBAT.")

    # -------------------------------------------------- Student's Declaration
    add_heading1(doc, "Student's Declaration")
    add_para(doc, "This is to certify that the work presented in this report titled “EduLearn — A Web-Based E-Learning Platform for Course Delivery, Student Progress Tracking, and Instructor Management” is the outcome of the investigation and development carried out by the following student under the supervision of Dr. Md. Abdul Awal, Associate Professor, Department of Computer Science and Engineering, International University of Business Agriculture and Technology.")
    add_spacer(doc, 24)
    add_para(doc, "__________________\n\nMd. Jobyadul Islam\nStudent ID 22303363", align="left", indent_first=False)

    # -------------------------------------------------- Supervisor's Certification
    add_heading1(doc, "Supervisor's Certification")
    add_para(doc, "This is to certify that the practicum report entitled “EduLearn — A Web-Based E-Learning Platform for Course Delivery, Student Progress Tracking, and Instructor Management” has been prepared by Md. Jobyadul Islam under my supervision. The work presented in this report is an outcome of the student's study, development, and practical implementation carried out as part of the practicum program.")
    add_para(doc, "I have reviewed the report and found it satisfactory for submission to the Practicum Defense Committee of the Department of Computer Science and Engineering, IUBAT—International University of Business Agriculture and Technology.")
    add_spacer(doc, 24)
    add_para(doc, "_______________________________\n\nDr. Md. Abdul Awal\nAssociate Professor\nDepartment of Computer Science and Engineering\nIUBAT—International University of Business Agriculture and Technology", align="left", indent_first=False)

    # -------------------------------------------------- Abstract
    add_heading1(doc, "Abstract")
    add_para(doc, "The rapid shift toward online and blended learning has created sustained demand for e-learning platforms that can manage course delivery, student progress, assessment, and payment in one coherent system rather than through disconnected spreadsheets, file shares, and third-party quiz tools. The objective of this Practicum is to design and develop EduLearn, a web-based e-learning platform that centralizes course creation, enrollment, lesson delivery, quiz and assignment assessment, certification, and platform administration for three distinct user roles.")
    add_para(doc, "The system is developed using ASP.NET Core MVC (.NET 10) with C# as the primary programming language, Microsoft SQL Server as the relational database, and Entity Framework Core (code-first, migration-driven) for data access. The system implements cookie-based authentication and role-based authorization for Student, Instructor, and Admin roles. Key features include course creation and a mandatory Admin approval-and-pricing workflow, module/lesson/assignment/quiz authoring with due-date enforcement, automatic quiz grading, enrollment and coupon-discounted checkout through the bKash payment gateway, lesson-progress tracking with PDF certificate issuance on completion, a Gmail-style read/unread notification system, PDF revenue and course-activity reporting, and an AI-assisted chat widget for course recommendations. A distinct, security-hardened instructor-onboarding flow — a Google OAuth or email access request, admin approval, a rate-limited one-time code, and only then the full application form — replaces what would otherwise be an insecure, manually-issued PIN.")
    add_para(doc, "The development process followed an incremental, Agile-influenced approach covering requirement engineering, system analysis, database design, implementation, and continuous testing, with a 33-case automated xUnit test suite exercising the system's core business rules (enrollment, certification eligibility, quiz grading, admin approval workflows, PDF generation) alongside extensive live, evidence-based manual verification of every feature against the running application and its database. This practicum provided practical experience in building a secure, role-based, real-world web application using modern .NET web technologies and structured software engineering practice, from initial requirements through to a system ready for production use.")

    # -------------------------------------------------- Acknowledgments
    add_heading1(doc, "Acknowledgments")
    add_para(doc, "First, I would like to thank the Almighty for His continuous blessings and guidance throughout my academic journey. Without His blessings, the completion of this work would not have been possible.")
    add_para(doc, "I am grateful to the late Professor Dr. Md. Alimullah Miyan, the founder and former Vice-Chancellor of the International University of Business, Agriculture and Technology (IUBAT), for establishing this institution and providing students with the opportunity to pursue higher education.")
    add_para(doc, "I would like to express my sincere gratitude to Professor Dr. Abdur Rab, Vice-Chancellor of IUBAT, for his valuable leadership and support.")
    add_para(doc, "I am also grateful to Prof. Dr. Utpal Kanti Das, Chairman and Professor, Department of Computer Science and Engineering, for his guidance and leadership.")
    add_para(doc, "My heartfelt thanks go to Shahinur Alam, Assistant Professor and Coordinator, Department of Computer Science and Engineering, for his continuous support and encouragement.")
    add_para(doc, "I would like to express my deepest gratitude to my respected supervisor, Dr. Md. Abdul Awal, Associate Professor, Department of Computer Science and Engineering, IUBAT, for his valuable guidance, suggestions, and continuous support throughout the development of EduLearn and the preparation of this report.")
    add_para(doc, "I am also thankful to Touch & Solve Technologies Ltd., and to Mr. Abul Kalam Azad, CEO, for providing a professional environment in which to complete my internship and apply my academic knowledge to real-world software development.")
    add_para(doc, "Finally, I sincerely thank my parents, friends, and classmates for their continuous encouragement, motivation, and support throughout this journey.")

    # -------------------------------------------------- ToC / LoF / LoT
    add_toc_page(doc)

    doc.add_page_break()
    p = doc.add_paragraph(style='Heading 1'); r = p.add_run("List of Figures"); set_run_font(r, 16, bold=True)
    figs = [
        "Figure 1.1 Agile Process Model Used for EduLearn Development.",
        "Figure 1.2 Technical Architecture and Technology Stack of EduLearn.",
        "Figure 2.1 Organizational Structure of Touch & Solve Technologies Ltd.",
        "Figure 3.1 Use Case Diagram of EduLearn.",
        "Figure 4.1 Activity Diagram of Admin Activities.",
        "Figure 4.2 Activity Diagram of Instructor Activities.",
        "Figure 4.3 Activity Diagram of Student Activities.",
        "Figure 4.4 Activity Diagram of the Instructor Access-Request and Onboarding Flow.",
        "Figure 4.5 Swim Lane Diagram of Course Approval and Enrollment.",
        "Figure 4.6 Sequence Diagram of Course Approval Flow.",
        "Figure 4.7 Sequence Diagram of Quiz Attempt and Auto-Grading Flow.",
        "Figure 7.1 System Architecture of EduLearn.",
        "Figure 7.2 Login Page.", "Figure 7.3 Admin Dashboard.", "Figure 7.4 Instructor Dashboard (My Courses).",
        "Figure 7.5 Student Dashboard (My Enrollments).", "Figure 7.6 Course Catalog and Course Details Page.",
        "Figure 7.7 Course Content Authoring Interface (Manage Content).", "Figure 7.8 Lesson Viewer with Assignment and Quiz.",
        "Figure 7.9 Quiz-Taking Interface.", "Figure 7.10 Cart and Checkout (bKash) Interface.",
        "Figure 7.11 Instructor Student Roster and Activity Report.", "Figure 7.12 Admin Course Overview (Instructor Details, Activity Timeline, Roster).",
        "Figure 7.13 Revenue Report Interface.", "Figure 7.14 Admin Coupon and Payment Management Interface.",
        "Figure 7.15 Context Level Diagram (Level 0).", "Figure 7.16 Level 1 Data Flow Diagram.",
        "Figure 7.17 Level 2 DFD — Authentication and Instructor Onboarding.",
        "Figure 7.18 Level 2 DFD — Course Creation and Approval.",
        "Figure 7.19 Level 2 DFD — Enrollment and Payment.",
        "Figure 7.20 Level 2 DFD — Lesson, Assignment, and Quiz Delivery.",
        "Figure 7.21 Level 2 DFD — Notification Processing.",
        "Figure 7.22 Level 2 DFD — Instructor and Admin Reporting.",
        "Figure 7.23 Level 2 DFD — Certificate and Receipt Generation.",
        "Figure 7.24 Entity Relationship Diagram of EduLearn.",
    ]
    for f in figs:
        add_para(doc, f, align="left", indent_first=False, space_after=4)

    doc.add_page_break()
    p = doc.add_paragraph(style='Heading 1'); r = p.add_run("List of Tables"); set_run_font(r, 16, bold=True)
    tabs = [
        "Table 1.1 Feasibility Summary.", "Table 3.1 Requirement Categories.", "Table 3.2 The major actor responsibilities.",
        "Table 5.1 Risk Analysis Table for EduLearn.", "Table 5.2 Risk Planning Table for EduLearn.",
        "Table 5.3 RMMM Plan No. 1.", "Table 5.4 RMMM Plan No. 2.", "Table 5.5 RMMM Plan No. 3.", "Table 5.6 RMMM Plan No. 4.",
        "Table 5.7 RMMM Plan No. 5.", "Table 5.8 RMMM Plan No. 6.", "Table 5.9 RMMM Plan No. 7.", "Table 5.10 RMMM Plan No. 8.",
        "Table 6.1 UFP for Transaction Functions.", "Table 6.2 UFP for Data Functions.", "Table 6.3 Total Degree of Influence.",
        "Table 6.4 Function Point Calculation Summary.", "Table 6.5 Development Resources.", "Table 6.6 Project Task Scheduling.",
        "Table 8.1 through Table 8.20 System Testing Design (one per major module).",
    ]
    for t in tabs:
        add_para(doc, t, align="left", indent_first=False, space_after=4)

    # ============================================================ CHAPTER 1
    add_heading1(doc, "Chapter 1. Introduction")
    add_para(doc, "Electronic learning (e-learning) has become a central part of modern education and workforce training, allowing institutions and independent instructors to deliver structured courses, assess learners, and track progress without requiring a shared physical classroom. An e-learning platform is a software system that lets an organization create courses, enroll and monitor learners, deliver graded assessments, and manage the commercial and administrative side of running those courses — enrollment, payment, and certification — through one centralized application. As both institutions and individual instructors increasingly depend on structured, trackable online delivery, an efficient e-learning platform can materially improve completion rates, administrative overhead, and the reliability of assessment.")
    add_para(doc, "In many small and mid-sized learning operations, course content, student progress, quiz grading, and payment are still managed through a patchwork of tools — a video host, a separate quiz tool, spreadsheets for enrollment and payment records, and email for everything else. This fragmentation makes it difficult to verify who has actually completed a course, to enforce deadlines consistently, to reconcile payments against access, and to give instructors a trustworthy picture of how their own students are doing.")
    add_para(doc, "To address these challenges, this practicum project focuses on the development of EduLearn, a web-based e-learning platform designed to centralize course delivery, assessment, enrollment, and payment in a single system with clearly separated responsibilities for Students, Instructors, and Admins. The system is developed using ASP.NET Core MVC (.NET 10) and C#, with Microsoft SQL Server managed through Entity Framework Core for all data access and schema evolution (via code-first migrations).")
    add_para(doc, "EduLearn provides course creation, structured content authoring (modules → lessons → assignments/quizzes), enrollment (free and paid), coupon-discounted checkout through the bKash payment gateway, automatic quiz grading with due-date enforcement, progress-based certificate issuance, student review submission, a Gmail-style notification system, and PDF-based reporting for both instructors and administrators. A deliberately security-conscious instructor-onboarding flow — Google OAuth or email access request, Admin approval, a rate-limited one-time verification code, and only then the actual application form — replaces what would otherwise be an easily-abused, manually-distributed PIN, and Admins retain sole authority over course pricing so instructors cannot misprice their own courses.")
    add_para(doc, "The primary goal of this project is to provide a structured, secure, and maintainable e-learning platform that improves course delivery, enforces meaningful deadlines, provides transparent progress and revenue reporting, and protects both learner data and the platform's payment flow. Development of this system also provided practical, hands-on experience applying software engineering concepts including requirement analysis, relational database design, MVC architecture, authentication and authorization, service-based business logic, automated and manual testing, and iterative, evidence-verified development.")

    add_heading2(doc, "1.1 Background of the Study")
    add_para(doc, "E-learning has evolved considerably alongside the broader growth of web application technology. In earlier eras, course delivery commonly relied on static content, shared files, and email-based communication between instructors and learners, with little to no centralized tracking of who had completed what, when a payment had actually cleared, or how a learner was progressing relative to a deadline. While workable at a very small scale, this approach breaks down as the number of courses, students, and instructors grows: progress cannot be verified reliably, deadlines are unenforced, and payment reconciliation becomes manual and error-prone.")
    add_para(doc, "With the growth of web-based application platforms, dedicated e-learning systems have become the standard way to organize course delivery. Modern platforms typically support course authoring, structured lesson delivery, automatic assessment, enrollment and payment processing, progress tracking, and reporting — allowing institutions and instructors to focus on content and pedagogy rather than on manual administration.")
    add_para(doc, "In a platform that supports paid courses specifically, correct enrollment-to-access mapping is particularly important: a student's access to full course content must correctly reflect whether they have paid, a coupon must be validated and applied consistently, and a refund must correctly revoke access rather than silently leaving it granted. Likewise, once assignments and quizzes carry real deadlines, the system must consistently prevent late submissions rather than only displaying a due date as decoration.")
    add_para(doc, "EduLearn is designed to address these operational requirements through a centralized, role-based web application. The system provides separate access and functionality for Student, Instructor, and Admin roles according to their responsibilities. It includes course authoring and a mandatory Admin approval-and-pricing gate, structured content delivery with enforced assignment/quiz deadlines, automatic quiz grading, enrollment and bKash-based paid checkout with coupon support, lesson-progress tracking with certificate issuance, student reviews, a notification system, and PDF-based revenue and course-activity reporting for both Instructors and Admins.")
    add_para(doc, "The system is developed using ASP.NET Core MVC (.NET 10) and C#, with Microsoft SQL Server and Entity Framework Core for database operations. The project follows a structured, service-based architecture — dedicated static service classes such as QuizGrader, CourseProgressCalculator, CourseRosterService, CourseReportService, and ReportPdfService keep business logic and PDF-generation logic out of the controllers — to maintain a scalable, testable, and maintainable application.")

    add_heading2(doc, "1.2 Methodology")
    add_para(doc, "EduLearn was developed using a systematic, structured approach that emphasizes maintainability, security, and correctness of business rules that directly affect money (payments, refunds, coupons) and academic integrity (deadlines, grading, certification). The system follows the Model-View-Controller (MVC) architectural pattern using ASP.NET Core MVC, where Controllers handle HTTP requests and orchestrate work, static Service classes contain the actual business logic and calculations, Models/ViewModels represent application and display data, and Razor Views (.cshtml) render the user interface server-side. Entity Framework Core (code-first) is used for all database access, with each schema change captured as an explicit, version-controlled migration.")
    add_para(doc, "The development process involved iterative requirement understanding, database and system design, feature implementation, and — distinctively for this project — a strict discipline of live verification: rather than trusting that a change “should” work, each feature was checked against the actual running application and its live SQL Server database (via direct queries and HTTP-level testing) before being considered complete, and again after any related fix. This discipline surfaced and corrected several real defects during development (see Chapter 5 and Chapter 8) that a purely code-review-based process would likely have missed.")
    add_para(doc, "Two types of data sources were used to complete this practicum work and develop this report: Primary Sources and Secondary Sources.")

    add_heading3(doc, "1.2.1 Primary Sources.")
    for b in ["Hands-on development of the EduLearn application using ASP.NET Core MVC and C#.",
              "Design and implementation of the database schema and its relationships using Entity Framework Core and SQL Server, evolved through more than a dozen incremental migrations.",
              "Development and testing of role-specific modules for Student, Instructor, and Admin.",
              "Implementation of course authoring, the Admin course-approval-and-pricing workflow, enrollment, checkout, and certification.",
              "Implementation of the Google OAuth/email instructor access-request flow, its rate-limited one-time verification code, and the downstream full application review.",
              "Development of quiz auto-grading, assignment/quiz due-date enforcement, student roster and weekly/monthly activity reporting, and PDF report/certificate/receipt generation.",
              "Manual, evidence-based testing (live HTTP requests and direct database verification) and automated xUnit testing, debugging, and refinement of system functionality throughout development."]:
        add_bullet(doc, b)

    add_heading3(doc, "1.2.2 Secondary Sources.")
    for b in ["Study of software engineering, relational database design, and web application development concepts.",
              "Review of official documentation for ASP.NET Core MVC, Entity Framework Core, and Microsoft SQL Server.",
              "Study of technical documentation related to Google OAuth 2.0 and the bKash payment gateway's tokenized checkout API.",
              "Review of resources related to authentication, authorization, email delivery (SMTP/MailKit), file upload handling, and PDF generation (QuestPDF).",
              "Reference to relevant online technical resources and developer documentation for troubleshooting and implementation guidance."]:
        add_bullet(doc, b)

    add_heading2(doc, "1.3 Objectives")
    add_para(doc, "The objectives of EduLearn have been divided into two categories: Broad Objective and Specific Objectives.")
    add_heading3(doc, "1.3.1 Broad Objective.")
    add_para(doc, "The primary objective of this practicum is to design and develop a secure, role-based, and maintainable e-learning platform that centralizes course authoring, enrollment, payment, assessment, progress tracking, certification, and reporting to improve the reliability and transparency of online course delivery.")
    add_heading3(doc, "1.3.2 Specific Objectives.")
    for b in ["To develop a secure authentication and authorization system with role-based access and separate permissions for Students, Instructors, and Admins.",
              "To manage the complete course lifecycle, including authoring, mandatory Admin approval and pricing, content structuring, and publication.",
              "To implement enrollment and payment processing — including free enrollment, coupon-discounted paid checkout via bKash, and Admin-issued refunds that correctly revoke access.",
              "To provide automatically-graded quizzes and deadline-enforced assignments, lesson-progress tracking, and certificate issuance on course completion.",
              "To provide Instructors and Admins with student-roster visibility and weekly/monthly activity reporting, exportable as PDF."]:
        add_bullet(doc, b)

    add_heading2(doc, "1.4 Process Model")
    add_para(doc, "An Agile, incrementally-driven process model was adopted for developing EduLearn. The system was built module by module — a working vertical slice of a feature (data model, migration, controller logic, and view) was completed, manually and automatically tested against the live application, and only then was the next feature layered on top. Several features (for example, the instructor-application flow, and course pricing authority) were revised more than once as live testing and direct user feedback surfaced gaps between the intended behavior and the actual behavior.")
    add_para(doc, "The rationale for selecting the Agile Process Model includes:")
    for b in ["Flexibility: requirements were refined as real usage revealed gaps — for example, discovering that Admin and Instructor accounts could not preview paid course content without an actual paid enrollment, and correcting this.",
              "Incremental Development: features were developed and integrated module by module rather than as one large, unverifiable batch.",
              "Continuous Testing: each module was verified live — via direct HTTP requests, database queries, and the automated xUnit suite — both when it was built and again whenever a related change was made.",
              "Early Error Detection: defects (see Chapter 5) were identified and corrected close to when they were introduced, rather than accumulating.",
              "Better Progress Monitoring: progress was evaluated feature by feature against real, observable behavior rather than against a static plan.",
              "Improved Maintainability: the modular, service-oriented approach kept the codebase organized as it grew."]:
        add_bullet(doc, b)
    add_figure(doc, "fig_1_1_agile.png", "Figure 1.1 Agile Process Model Used for EduLearn Development.")
    add_para(doc, "Figure 1.1 illustrates the Agile-based development process followed for EduLearn. Development activities were performed progressively, with each module implemented, verified against the live system, and refined before the next was layered on top.")

    add_heading2(doc, "1.5 Feasibility Study")
    add_para(doc, "Before developing EduLearn, a feasibility study was conducted to determine whether the project was practical and achievable using the available technical, financial, and operational resources. The feasibility of EduLearn was analyzed from three perspectives: technical feasibility, economic feasibility, and operational feasibility.")
    add_heading3(doc, "1.5.1 Technical Feasibility.")
    add_para(doc, "EduLearn is technically feasible because it is developed using established, well-documented, and widely used web development technologies. The application is built using ASP.NET Core MVC (.NET 10) and C#, with Entity Framework Core for database operations and Microsoft SQL Server as the relational database engine.")
    add_para(doc, "The system also uses Google OAuth 2.0 for the instructor-onboarding access-request flow, MailKit-based SMTP email delivery for OTP codes and workflow notifications, the bKash tokenized checkout API for payment processing, QuestPDF for certificate/receipt/report generation, Chart.js for administrative and reporting charts, and standard web technologies — HTML, CSS, Bootstrap 5, and JavaScript — for the user interface. All of these were successfully integrated and verified working during this practicum.")
    add_figure(doc, "fig_1_2_stack.png", "Figure 1.2 Technical Architecture and Technology Stack of EduLearn.")
    add_para(doc, "Figure 1.2 presents the major technologies and supporting services used to develop and operate EduLearn, combining ASP.NET Core MVC, C#, Entity Framework Core, and SQL Server with frontend technologies and external services such as Google OAuth 2.0, SMTP-based email, the bKash payment gateway, and QuestPDF.")
    add_heading3(doc, "1.5.2 Economic Feasibility.")
    add_para(doc, "EduLearn is economically feasible because the development environment primarily uses freely available and open-source technologies. ASP.NET Core, Entity Framework Core, Bootstrap, and the xUnit testing framework can be used for development without licensing costs; SQL Server Express is likewise free for development and small-scale deployment.")
    add_para(doc, "The system can also reduce operational dependency on manual, spreadsheet-based enrollment and payment tracking by centralizing course delivery, enrollment, payment, assessment, and reporting within a single platform, providing significant operational benefit without requiring substantial upfront investment in proprietary software.")
    add_heading3(doc, "1.5.3 Operational Feasibility.")
    add_para(doc, "EduLearn is operationally feasible because it is designed around the concrete responsibilities of its three user roles. The system provides separate access and permissions for Student, Instructor, and Admin, so each user is only ever presented with the functionality relevant to their responsibilities — for example, an Instructor cannot set their own course's price, and only an Admin can issue a refund.")
    add_para(doc, "The system centralizes course content, enrollment and payment records, assessment results, notifications, and reports. Automatic quiz grading, deadline enforcement, and PDF report generation reduce manual effort and improve the reliability of monitoring course activity. Therefore, the proposed system can support real course-delivery and administrative operations and is operationally feasible.")
    add_table(doc, [
        ["Feasibility Type", "Description", "Conclusion"],
        ["Technical Feasibility", "The system uses available, well-documented technologies: ASP.NET Core MVC, C#, Entity Framework Core, SQL Server, Google OAuth 2.0, bKash, MailKit, and QuestPDF.", "Technically Feasible"],
        ["Economic Feasibility", "The system primarily uses freely available development technologies and reduces dependency on manual enrollment/payment tracking.", "Economically Feasible"],
        ["Operational Feasibility", "Role-based access, centralized course/enrollment/payment management, and automated reporting support real course-delivery operations.", "Operationally Feasible"],
    ], caption="Table 1.1 Feasibility Summary.", col_widths=[1.45, 3.19, 1.26])
    add_para(doc, "Table 1.1 summarizes the feasibility of EduLearn from technical, economic, and operational perspectives. The analysis indicates that the system is feasible using the available resources and technologies — and, unlike a purely theoretical feasibility study, this is additionally confirmed by the fact that the system was actually built, and every feature described in this report was verified working against a live running instance.")

    add_heading2(doc, "1.6 Structure of the Report")
    add_para(doc, "The report is organized into ten chapters, each focusing on a specific aspect of the practicum project. The chapters are arranged to present the development process progressively, from the introduction and organizational background to requirements, analysis, project management, planning, system design, testing, ethical considerations, and conclusion.")
    for b in ["Chapter 1 – Introduction: Presents the background of the study, methodology, objectives, process model, feasibility study, and structure of the report.",
              "Chapter 2 – Organizational Overview: Describes Touch & Solve Technologies Ltd., including its vision, services, organizational structure, and the student's role within it.",
              "Chapter 3 – Requirement Engineering: Presents the user, system, functional, and non-functional requirements of EduLearn.",
              "Chapter 4 – Analysis: Describes the system analysis using use case, activity, swimlane, and sequence models.",
              "Chapter 5 – Project Management: Discusses risk identification, analysis, planning, and RMMM strategies, grounded in issues actually encountered during development.",
              "Chapter 6 – Project Planning: Presents Function Point-based effort estimation, resource planning, and the development schedule.",
              "Chapter 7 – System Design: Describes the system architecture, internal interface design, data flow, entity relationships, and database design.",
              "Chapter 8 – Quality Assurance and Testing: Presents the testing methodology, test cases, and quality assurance activities performed to verify EduLearn's functionality and reliability.",
              "Chapter 9 – Ethical Considerations: Discusses privacy, security, fairness, and sustainability considerations relevant to a platform handling student data and payments.",
              "Chapter 10 – Conclusion: Summarizes the project outcomes, major contributions, limitations, and possible future improvements."]:
        add_bullet(doc, b)

    # ============================================================ CHAPTER 2
    add_heading1(doc, "Chapter 2. Organizational Overview")
    add_para(doc, "Touch & Solve Technologies Ltd. is an IT company based in Dhaka, Bangladesh, that began its journey in 2009 with customized software development and expanded by 2011 into broader IT-enabled services, positioning itself as a system integrator for data centers, networking, surveillance, and IT solutions for government and non-government institutions. The company is headquartered at Plot 6 (3rd & 6th Floor), Road 1/A, Sector 17, Uttara, Dhaka 1230, Bangladesh.")
    add_para(doc, "The company builds standard, fully responsive website development services along with mobile and desktop application development, and additionally offers ICT solutions, IT infrastructure and LAN-Wifi services, customized software (POS, Institute Management, Hospital Management), networking solutions, CCTV, and time-attendance systems. Over the years, the company has served a large number of corporate and institutional clients across Bangladesh.")

    add_heading2(doc, "2.1 Organization Vision")
    add_para(doc, "Touch & Solve Technologies Ltd. positions its work within the broader vision of a smart, digitally-enabled Bangladesh — helping government and private institutions modernize their operations through well-designed, dependable software and IT infrastructure.")

    add_heading2(doc, "2.2 Organization Mission")
    add_para(doc, "The company's mission is to design, develop, and support standard and customized software solutions — from responsive websites to institute and hospital management systems — while building long-term technical partnerships with its clients through reliable IT infrastructure, networking, and surveillance services.")

    add_heading2(doc, "2.3 Organization Services")
    for b in ["Website Design and Development — responsive, fully custom website development for corporate and institutional clients.",
              "Web and Application Development — customized web applications and mobile/desktop application development.",
              "Software Development Solutions — customized POS, Institute Management, and Hospital Management software.",
              "IT Infrastructure and Networking — LAN/Wi-Fi design and deployment, data center and campus networking.",
              "Surveillance and Security Systems — CCTV and time-attendance solutions.",
              "ICT Solutions for Government and Non-Government Institutions — system integration services for public and private sector clients."]:
        add_bullet(doc, b)

    add_heading2(doc, "2.4 Organizational Structure")
    add_para(doc, "The organogram of the organization is shown in Figure 2.1.")
    add_figure(doc, "fig_2_1_org_structure.png", "Figure 2.1 Organizational Structure of Touch & Solve Technologies Ltd.")

    add_heading2(doc, "2.5 My Position in this Organization")
    add_para(doc, "At Touch & Solve Technologies Ltd., I worked as a Web Development Intern. My role focused on the design, development, and testing of web-based applications under the guidance of the company's development team. During my internship, I was engaged in full-stack web application development, which directly informed and complemented the independent design and development of the EduLearn platform described in this report.")
    add_para(doc, "My responsibilities included assisting with web application development tasks, applying ASP.NET Core MVC and C# in practical assignments, working with relational databases, and gaining exposure to real client-facing software development workflows. Internship period: [Start Date] to [End Date].")
    add_para(doc, "This position provided valuable hands-on exposure to professional software development practice and reinforced the technical and problem-solving skills applied throughout the independent development of EduLearn.")

    add_heading2(doc, "2.6 Address of the Organization")
    add_para(doc, "Touch & Solve Technologies Ltd.\nPlot 6 (3rd & 6th Floor), Road 1/A, Sector 17, Uttara, Dhaka 1230, Bangladesh.\nPhone: +880 1958227203 / +880 1958227200\nEmail: info@touchandsolve.com\nWebsite: www.touchandsolve.com", align="left", indent_first=False)

    # ============================================================ CHAPTER 3
    add_heading1(doc, "Chapter 3. Requirement Engineering")
    add_para(doc, "Requirement Engineering is the systematic process of identifying, analyzing, documenting, validating, and managing the requirements of a software system. It establishes a clear understanding of the functions the system must provide, the users who will interact with it, and the constraints under which it must operate. Proper requirement engineering helps reduce ambiguity, control project scope, and provides a strong foundation for system analysis, design, implementation, and testing.")
    add_para(doc, "For EduLearn, requirement engineering was carried out by analyzing the operational workflow of course delivery, enrollment and payment, assessment, instructor onboarding, and platform administration. Requirements were refined throughout development as live testing against the running application revealed gaps between intended and actual behavior — most notably around course pricing authority, instructor-application security, and Admin/Instructor content-preview access, each of which prompted a requirement correction that is reflected in the functional requirements below.")

    add_heading2(doc, "3.1 Requirement Analysis")
    add_para(doc, "The requirements of EduLearn are organized into three major categories: User and System Requirements, Functional Requirements, and Non-Functional Requirements.")
    add_table(doc, [
        ["Requirement Category", "Description"],
        ["User and System Requirements", "Defines the system users, roles, access privileges, and major system expectations."],
        ["Functional Requirements", "Defines the specific operations and services provided by EduLearn."],
        ["Non-Functional Requirements", "Defines security, performance, reliability, usability, scalability, maintainability, and data integrity requirements."],
    ], caption="Table 3.1 Requirement Categories.", col_widths=[1.93, 3.97])

    add_heading3(doc, "3.1.1 User and System Requirements")
    add_para(doc, "EduLearn supports three roles with different responsibilities and access privileges. The system must ensure that each user can access only the functions permitted for their role.")
    for b in [
        "Users must be able to securely register, log in, and log out of EduLearn. Self-registration verifies the new account's email through a one-time code before the account is created, and only ever creates Student accounts — becoming an Instructor requires the separate access-request pipeline. Passwords must be securely hashed before storage, and the system must maintain authenticated sessions using cookie-based authentication.",
        "The system must support role-based access control for Student, Instructor, and Admin. Each role must have separate permissions, navigation, and access restrictions, enforced via [Authorize(Roles = \"...\")] on controllers/actions, not merely hidden in the UI. Admin must have the highest level of system management authority, including sole authority over course pricing.",
        "Administrators must be able to manage system users and monitor platform activity, including reviewing and approving/denying instructor access requests and full applications, and maintaining an IsRejected/IsActive status so deactivated or rejected accounts are excluded from active-user counts and lists consistently.",
        "The system must provide centralized course management, supporting course creation, editing, and deletion, gated by an Admin approval workflow that also sets the course's price. Any edit to an already-approved course must return it to Pending for re-review before the change goes live.",
        "The system must support structured course content and enforced deadlines. Assignments and Quizzes must carry a due date and time, and the system must reject (server-side, not just hide the button for) any submission or attempt made after that deadline has passed. Quizzes must be automatically graded against their correct-option data at submission time.",
        "The system must support a secure, self-service instructor-onboarding flow: a visitor requests instructor access via Google OAuth or a typed email address, an Admin must approve the request, and approval emails a unique link and a 6-digit one-time code that must expire and be rate-limited before the applicant can reach the actual application form.",
    ]:
        add_bullet(doc, b)

    add_heading3(doc, "3.1.2 Functional Requirements")
    add_para(doc, "The functional requirements describe the specific operations and services that EduLearn must provide to its users.")
    for h, items in [
        ("User Authentication and Authorization", [
            "Users must be able to register (Student, with OTP-verified email) and log in using valid credentials.",
            "The system must verify passwords using ASP.NET Core Identity's secure password hashing.",
            "Unauthorized users must not be able to access protected Instructor or Admin controllers/actions, independent of what the navigation menu shows.",
            "Users must be redirected to a role-appropriate landing page after login, and deactivated users must not be permitted to access the system.",
        ]),
        ("Instructor Access Request and Onboarding", [
            "A request may be created via Google OAuth or a typed, validated email address; a second request from the same still-pending email must not create a duplicate row.",
            "Admin approval must generate a cryptographically random one-time code and access token, and must email both together.",
            "The code-entry action must independently verify the request is Approved, not already consumed, not expired, and under the maximum attempt count — on every attempt, not only the first.",
            "Only after successful code verification does the session gain access to the actual application form.",
        ]),
        ("Course Management", [
            "Instructors must be able to create, edit, and delete their own courses, but must not be able to set or change their own course's price.",
            "Admin must be able to approve a Pending course, choosing at that moment whether it is Free or setting a fixed price, or reject it with an optional reason emailed back to the Instructor.",
            "Admin must be able to edit the price of, or permanently delete, any course at any time from Manage Courses.",
        ]),
        ("Content Authoring", [
            "Instructors must be able to add Modules, Lessons, Assignments, and Quizzes under their own courses.",
            "Assignment and Quiz creation forms must accept both a date and a time for the due date.",
            "Quiz authoring must accept a title, multiple-choice questions with marked correct options, a pass-mark percentage, a time limit, and a due date.",
        ]),
        ("Enrollment and Payment", [
            "Free-course enrollment must be created immediately as Active; paid-course enrollment must be created as Pending in the student's Cart.",
            "The Cart must allow applying one coupon code, computing and displaying the discounted total before checkout.",
            "Checkout must be processed through the bKash gateway; on a successful callback, the enrollment must flip to Active and a Payment record created.",
            "Admin must be able to refund a successful payment, reverting the enrollment to Pending and notifying the student.",
        ]),
        ("Lesson Progress, Quizzes, and Assignments", [
            "Students must be able to mark a lesson complete; progress percentage must be computed as completed lessons ÷ total lessons for that course.",
            "Quiz submissions must be graded automatically by QuizGrader against stored correct answers, updating in place on retake rather than accumulating duplicate rows.",
            "Both quiz attempts and assignment submissions must be rejected — server-side, on every entry point — once the item's due date has passed.",
        ]),
        ("Certification, Reviews, Notifications, Reporting", [
            "A certificate must only be downloadable once every lesson in the course is marked complete for that student.",
            "Students must be able to write a review only after completing a course, and edit it afterward; one review per student per course.",
            "The system must create a Notification row for events including course approval/rejection, new enrollment, and instructor onboarding decisions, visibly distinguished as unread versus read.",
            "Instructors and Admin must be able to view a student roster and generate weekly/monthly activity reports, exportable as PDF; Admin additionally has platform-wide Revenue and Course Analytics reports.",
        ]),
        ("Admin/Instructor Content Preview", [
            "Admin and Instructor accounts must be able to open and read any course's lessons, assignments, and quizzes — including paid courses they have not purchased — without an enrollment record being created.",
            "Because they are previewing, not taking the course, the interactive “Mark as Complete,” “Submit Assignment,” and “Take Quiz” controls must not be shown to them on those pages.",
        ]),
    ]:
        add_heading3(doc, h)
        for it in items:
            add_bullet(doc, it)

    add_heading3(doc, "3.1.3 Non-Functional Requirements")
    for h, items in [
        ("Security", ["The system must authenticate users before allowing access to protected modules, and passwords must never be stored as plain text.",
                      "Role-based authorization must restrict access to role-specific controllers and actions at the server, not only via hidden UI elements.",
                      "The instructor one-time code must be rate-limited and time-limited so it cannot be brute-forced even by someone who has obtained the access-token link.",
                      "Google OAuth must be used only to read a verified email and must never create a persistent site login on its own.",
                      "Uploaded files must be validated against allowed extensions and size limits before being persisted."]),
        ("Performance", ["The system should remain responsive during normal interactive use.",
                         "Roster and activity-report queries must be built to avoid N+1 query patterns.",
                         "Pagination must be used for course listings and other potentially large result sets."]),
        ("Reliability", ["Enrollment, payment, and refund state must remain consistent — a refund must always revoke the access it granted.",
                         "Quiz retakes must update the existing result row in place rather than creating duplicate history.",
                         "A course edited by its Instructor must always return to Pending, never silently stay Approved with unreviewed changes live."]),
        ("Usability", ["Each role must have a dashboard and navigation appropriate to its responsibilities.",
                       "Deadlines, statuses, and validation errors must be shown in plain, specific language.",
                       "Notifications must make unread items visually obvious at a glance."]),
        ("Scalability and Maintainability", ["The application must follow the MVC architecture with business/report/calculation logic factored into static Service classes rather than left inside controllers.",
                                             "ViewModels must be used to separate view-facing data shapes from database entities where the two diverge.",
                                             "New features should be addable through new Services/Controllers/Views following the existing pattern without restructuring what already exists."]),
        ("Data Integrity", ["Primary and foreign keys must maintain valid relationships between Users, Courses, Enrollments, Payments, Quizzes, Assignments, and their submissions/results.",
                            "A rejected instructor application's account is deleted, but its details are preserved in a standalone archive table for audit purposes.",
                            "Required fields must be validated server-side, not only client-side."]),
    ]:
        add_heading3(doc, h)
        for it in items:
            add_bullet(doc, it)

    add_heading2(doc, "3.2 Use Case Diagram of the System")
    add_heading3(doc, "3.2.1 Use Case Diagram")
    add_para(doc, "The Use Case Diagram provides a high-level representation of EduLearn's functional scope. It identifies the major actors — Student, Instructor, and Admin — and illustrates the principal functions available to each, focusing on what each actor can do rather than how it is implemented internally.")
    add_table(doc, [
        ["Actor", "Major Use Cases"],
        ["Student", "Register/Login, Browse Courses, Enroll (Free/Paid), Cart & Coupon, Checkout, View Lessons, Mark Complete, Take Quiz, Submit Assignment, Write/Edit Review, Download Certificate/Receipt, View Notifications, Chat with Assistant"],
        ["Instructor", "Login, Create/Edit/Delete Course, Author Modules/Lessons/Assignments/Quizzes, View Reviews & Reply, View Quiz Results, View Student Roster, View/Export Activity Report"],
        ["Admin", "Login, Manage Users, Manage Categories, Approve/Reject Course (set Price), Manage Courses (Edit Price/Delete), Approve/Deny Access Request, Approve/Reject Full Application, View Course Overview (any course), View/Export Revenue & Analytics Reports, Manage Payments/Refunds, Manage Coupons"],
    ], caption="Table 3.2 The major actor responsibilities.", col_widths=[1.26, 4.64])
    add_figure(doc, "fig_3_1_usecase.png", "Figure 3.1 Use Case Diagram of EduLearn.", width_in=5.8)
    add_para(doc, "Figure 3.1 illustrates the major actors and their interactions with EduLearn, including authentication, course management, enrollment and payment, assessment, and administration.")
