"""Tiny rendering engine: turns a list of simple content tuples into a fully
styled .docx (Times New Roman, justified body text, numbered headings, real
Word TOC/page-number fields, tables, and embedded figures)."""

import os
from docx import Document
from docx.shared import Pt, Inches, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from docx.enum.section import WD_SECTION

FIG_DIR = os.path.join(os.path.dirname(__file__), "figures")
FONT = "Times New Roman"


def set_run_font(run, size=12, bold=False, italic=False, color=None):
    run.font.name = FONT
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    if color:
        run.font.color.rgb = RGBColor(*color)
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.find(qn('w:rFonts'))
    if rFonts is None:
        rFonts = OxmlElement('w:rFonts')
        rPr.append(rFonts)
    rFonts.set(qn('w:eastAsia'), FONT)


def add_field(paragraph, field_code, result_text=""):
    """Insert a real Word field (e.g. TOC, PAGE) so Word populates it on open/update."""
    run = paragraph.add_run()
    fld_begin = OxmlElement('w:fldChar')
    fld_begin.set(qn('w:fldCharType'), 'begin')
    instr = OxmlElement('w:instrText')
    instr.set(qn('xml:space'), 'preserve')
    instr.text = field_code
    fld_sep = OxmlElement('w:fldChar')
    fld_sep.set(qn('w:fldCharType'), 'separate')
    fld_text = OxmlElement('w:t')
    fld_text.text = result_text
    fld_end = OxmlElement('w:fldChar')
    fld_end.set(qn('w:fldCharType'), 'end')
    r = run._element
    r.append(fld_begin)
    r.append(instr)
    r.append(fld_sep)
    t_run = OxmlElement('w:r')
    t_run.append(fld_text)
    r.addnext(t_run)
    t_run.addnext(fld_end)
    set_run_font(run, 12)


def new_document():
    doc = Document()
    section = doc.sections[0]
    section.page_height = Inches(11.69)
    section.page_width = Inches(8.27)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1.25)
    section.right_margin = Inches(1)

    normal = doc.styles['Normal']
    normal.font.name = FONT
    normal.font.size = Pt(12)
    normal.paragraph_format.line_spacing = 1.5
    normal.paragraph_format.space_after = Pt(10)
    rpr = normal.element.get_or_add_rPr()
    rFonts = rpr.find(qn('w:rFonts'))
    if rFonts is None:
        rFonts = OxmlElement('w:rFonts')
        rpr.append(rFonts)
    rFonts.set(qn('w:eastAsia'), FONT)

    # Built-in Heading 1/2/3 styles, restyled to Times New Roman/black so the
    # native Word TOC field (which scans these styles) works out of the box.
    for name, size, center in (('Heading 1', 16, True), ('Heading 2', 14, False), ('Heading 3', 12.5, False)):
        st = doc.styles[name]
        st.font.name = FONT
        st.font.size = Pt(size)
        st.font.bold = True
        st.font.color.rgb = RGBColor(0, 0, 0)
        st.font.italic = (name == 'Heading 3')
        st.paragraph_format.space_before = Pt(18 if name == 'Heading 1' else (14 if name == 'Heading 2' else 10))
        st.paragraph_format.space_after = Pt(10 if name == 'Heading 1' else 6)
        st.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER if center else WD_ALIGN_PARAGRAPH.LEFT
        st.paragraph_format.line_spacing = 1.0
        rpr = st.element.get_or_add_rPr()
        rFonts = rpr.find(qn('w:rFonts'))
        if rFonts is None:
            rFonts = OxmlElement('w:rFonts')
            rpr.append(rFonts)
        rFonts.set(qn('w:eastAsia'), FONT)

    # footer page number, centered
    footer = section.footer
    fp = footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_field(fp, "PAGE")

    return doc


def add_heading1(doc, text, page_break=True):
    if page_break:
        doc.add_page_break()
    p = doc.add_paragraph(style='Heading 1')
    run = p.add_run(text)
    set_run_font(run, 16, bold=True)
    return p


def add_heading2(doc, text):
    p = doc.add_paragraph(style='Heading 2')
    run = p.add_run(text)
    set_run_font(run, 14, bold=True)
    return p


def add_heading3(doc, text):
    p = doc.add_paragraph(style='Heading 3')
    run = p.add_run(text)
    set_run_font(run, 12.5, bold=True, italic=True)
    return p


def add_para(doc, text, align="justify", size=12, italic=False, bold=False, space_after=10, indent_first=True):
    p = doc.add_paragraph()
    align_map = {"justify": WD_ALIGN_PARAGRAPH.JUSTIFY, "center": WD_ALIGN_PARAGRAPH.CENTER,
                 "left": WD_ALIGN_PARAGRAPH.LEFT, "right": WD_ALIGN_PARAGRAPH.RIGHT}
    p.alignment = align_map.get(align, WD_ALIGN_PARAGRAPH.JUSTIFY)
    p.paragraph_format.space_after = Pt(space_after)
    if indent_first and align == "justify":
        p.paragraph_format.first_line_indent = Inches(0.3)
    run = p.add_run(text)
    set_run_font(run, size, bold=bold, italic=italic)
    return p


def add_bullet(doc, text, size=12):
    p = doc.add_paragraph(style='List Bullet')
    p.paragraph_format.space_after = Pt(4)
    run = p.add_run(text)
    set_run_font(run, size)
    return p


def add_spacer(doc, pts=8):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(pts)


def add_figure(doc, filename, caption, width_in=5.6):
    path = os.path.join(FIG_DIR, filename)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run()
    if os.path.exists(path):
        run.add_picture(path, width=Inches(width_in))
    cap = doc.add_paragraph()
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap.paragraph_format.space_after = Pt(14)
    crun = cap.add_run(caption)
    set_run_font(crun, 10.5, bold=True)


def add_screenshot_placeholder(doc, caption, instruction):
    """Bordered box marking where a real app screenshot goes, with capture instructions."""
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = table.rows[0].cells[0]
    cell.width = Inches(5.6)
    tcPr = cell._tc.get_or_add_tcPr()
    borders = OxmlElement('w:tcBorders')
    for edge in ('top', 'left', 'bottom', 'right'):
        el = OxmlElement(f'w:{edge}')
        el.set(qn('w:val'), 'dashed')
        el.set(qn('w:sz'), '10')
        el.set(qn('w:color'), '888888')
        borders.append(el)
    tcPr.append(borders)
    shd = OxmlElement('w:shd')
    shd.set(qn('w:fill'), 'F5F5F5')
    tcPr.append(shd)
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(30)
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run("[ SCREENSHOT TO BE INSERTED HERE ]")
    set_run_font(r, 11, bold=True, color=(120, 120, 120))
    p2 = cell.add_paragraph()
    p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p2.paragraph_format.space_after = Pt(30)
    r2 = p2.add_run(instruction)
    set_run_font(r2, 9, italic=True, color=(120, 120, 120))
    cap = doc.add_paragraph()
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap.paragraph_format.space_after = Pt(14)
    crun = cap.add_run(caption)
    set_run_font(crun, 10.5, bold=True)


def _set_cell_border(cell, color="A0A0A0"):
    tcPr = cell._tc.get_or_add_tcPr()
    borders = OxmlElement('w:tcBorders')
    for edge in ('top', 'left', 'bottom', 'right'):
        el = OxmlElement(f'w:{edge}')
        el.set(qn('w:val'), 'single')
        el.set(qn('w:sz'), '4')
        el.set(qn('w:color'), color)
        borders.append(el)
    tcPr.append(borders)


def add_table(doc, rows, caption=None, header=True, col_widths=None, size=10, header_fill="D9D9D9"):
    if caption:
        cap = doc.add_paragraph()
        cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        cap.paragraph_format.space_before = Pt(10)
        cap.paragraph_format.space_after = Pt(6)
        crun = cap.add_run(caption)
        set_run_font(crun, 10.5, bold=True)

    n_cols = len(rows[0])
    table = doc.add_table(rows=0, cols=n_cols)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    for ridx, row in enumerate(rows):
        cells = table.add_row().cells
        is_header = header and ridx == 0
        for cidx, val in enumerate(row):
            cell = cells[cidx]
            _set_cell_border(cell)
            if is_header:
                shd = OxmlElement('w:shd')
                shd.set(qn('w:fill'), header_fill)
                cell._tc.get_or_add_tcPr().append(shd)
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.space_after = Pt(2)
            r = p.add_run(str(val))
            set_run_font(r, size, bold=is_header)
    if col_widths:
        for i, w in enumerate(col_widths):
            for row in table.rows:
                row.cells[i].width = Inches(w)
    add_spacer(doc, 10)
    return table


def add_toc_page(doc, title="Table of Contents"):
    doc.add_page_break()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(title)
    set_run_font(run, 16, bold=True)
    p.paragraph_format.space_after = Pt(18)
    p2 = doc.add_paragraph()
    add_field(p2, 'TOC \\o "1-3" \\h \\z \\u',
              "Right-click here and choose “Update Field” (or press F9) to generate the Table of Contents.")
    note = doc.add_paragraph()
    r = note.add_run("(Right-click above → Update Field, after opening this document in Word, to populate the Table of Contents with real page numbers.)")
    set_run_font(r, 9, italic=True, color=(150, 90, 20))
